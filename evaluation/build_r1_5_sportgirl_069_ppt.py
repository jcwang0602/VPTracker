#!/usr/bin/env python3
"""Export the R1.5 similar-athlete case and create an editable 069 PPT figure."""

import csv
import shutil
from pathlib import Path

import cv2
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = Path("/mnt/shared-storage-user/mineru4s/jcwang/VPTrack/data/TNLLT/JE_SportGirl_video_07")
PREDICTION_PATH = ROOT / (
    "outputs_vpt_r1/results_batchpool/"
    "069_qwen35_2b_vp_blue_w1_op100_out25_pixelcoords_no_imgtok_cap_1m_lr2e-5_bs128_e1_pdbs4_ga4_pool/tnllt/"
    "JE_SportGirl_video_07.txt"
)
OUTPUT_ROOT = ROOT / "paper/revision_results_qwen35_2b/R1_5_similar_targets_069"
PPT_PATH = OUTPUT_ROOT / "R1_5_similar_targets_069.pptx"
FIGURE_PATHS = [ROOT / "paper/VPTracker_PRL/sim_vis_vp_01.png"]
FRAMES = [45, 46, 47, 48]
FONT_PATH = "/usr/share/fonts/opentype/urw-base35/NimbusRoman-Regular.otf"

# BGR colors for OpenCV / RGB values in PowerPoint are converted below.
GT_COLOR = (0, 220, 0)
VP_COLOR = (255, 0, 0)
PREDICTION_COLOR = (0, 165, 255)


def load_xywh(path):
    return np.atleast_2d(np.loadtxt(path, delimiter=","))[:, :4]


def rgb(bgr):
    return RGBColor(bgr[2], bgr[1], bgr[0])


def vp_box(box, image_width, image_height):
    x, y, width, height = box
    center_x, center_y = x + width / 2, y + height / 2
    width, height = width * 4, height * 4
    left = max(0.0, min(image_width, center_x - width / 2))
    top = max(0.0, min(image_height, center_y - height / 2))
    right = max(0.0, min(image_width, center_x + width / 2))
    bottom = max(0.0, min(image_height, center_y + height / 2))
    return np.array([left, top, max(0.0, right - left), max(0.0, bottom - top)])


def draw_box(image, box, color, thickness):
    x, y, width, height = np.rint(box).astype(int)
    if width > 0 and height > 0:
        cv2.rectangle(image, (x, y), (x + width - 1, y + height - 1), color, thickness, cv2.LINE_AA)


def add_text(slide, text, x, y, width, height, size, color=RGBColor(0, 0, 0), align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    return shape


def add_outline(slide, box, image_width, image_height, x, y, width, height, color, line_width):
    left, top, box_width, box_height = box
    if box_width <= 0 or box_height <= 0:
        return
    rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x + left / image_width * width),
        Inches(y + top / image_height * height),
        Inches(box_width / image_width * width),
        Inches(box_height / image_height * height),
    )
    rect.fill.background()
    rect.line.color.rgb = rgb(color)
    rect.line.width = Pt(line_width)


def text_width(draw, text, font):
    return draw.textbbox((0, 0), text, font=font)[2]


def main():
    frames_root = OUTPUT_ROOT / "frames"
    inputs_root = OUTPUT_ROOT / "inputs"
    rendered_root = OUTPUT_ROOT / "rendered"
    for directory in (frames_root, inputs_root, rendered_root):
        directory.mkdir(parents=True, exist_ok=True)

    image_paths = sorted((DATA_ROOT / "imgs").iterdir())
    gt = load_xywh(DATA_ROOT / "groundtruth.txt")
    predictions = load_xywh(PREDICTION_PATH)
    if not (len(image_paths) == len(gt) == len(predictions)):
        raise ValueError("TNLLT image, GT, and prediction lengths differ")
    shutil.copy2(DATA_ROOT / "groundtruth.txt", inputs_root / "JE_SportGirl_video_07__GT.txt")
    shutil.copy2(PREDICTION_PATH, inputs_root / "JE_SportGirl_video_07__VPTracker_069.txt")
    language_source = DATA_ROOT / "language.txt"
    if language_source.exists():
        shutil.copy2(language_source, inputs_root / "language.txt")

    rows = []
    rendered = []
    for frame_number in FRAMES:
        index = frame_number - 1
        source = image_paths[index]
        frame_copy = frames_root / f"frame_{frame_number:04d}{source.suffix.lower()}"
        shutil.copy2(source, frame_copy)
        image = cv2.imread(str(source))
        height, width = image.shape[:2]
        prompt = vp_box(predictions[index - 1], width, height) if index else np.zeros(4)
        draw_box(image, prompt, VP_COLOR, 1)
        draw_box(image, gt[index], GT_COLOR, 3)
        draw_box(image, predictions[index], PREDICTION_COLOR, 3)
        rendered_path = rendered_root / f"frame_{frame_number:04d}.png"
        cv2.imwrite(str(rendered_path), image)
        rendered.append(rendered_path)
        for label, box in (("Ground Truth", gt[index]), ("Visual Prompt", prompt), ("VPTracker (069)", predictions[index])):
            rows.append({
                "frame": frame_number, "element": label,
                "x": f"{box[0]:.6f}", "y": f"{box[1]:.6f}",
                "width": f"{box[2]:.6f}", "height": f"{box[3]:.6f}",
                "source_image": str(frame_copy.relative_to(OUTPUT_ROOT)),
            })

    with (OUTPUT_ROOT / "bounding_boxes.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    (OUTPUT_ROOT / "README.md").write_text(
        "# R1.5 similar-target qualitative case from experiment 069\n\n"
        "The figure uses TNLLT `JE_SportGirl_video_07`, frames 0045--0048. "
        "The visual prompt is reconstructed from the preceding 069 prediction with scale 4, blue, 1 px, solid, and opacity 1.0.\n\n"
        "- `R1_5_similar_targets_069.pptx`: editable PowerPoint figure.\n"
        "- `frames/`: original source frames.\n"
        "- `rendered/`: raster reference render.\n"
        "- `inputs/`: complete GT and 069 prediction files.\n"
        "- `bounding_boxes.csv`: all displayed coordinates.\n"
    )

    panel_width, panel_height, gutter = 1180, 788, 40
    canvas_width = panel_width * 4 + gutter * 3
    canvas_height = panel_height + 220
    canvas = np.full((canvas_height, canvas_width, 3), 255, dtype=np.uint8)
    for column, frame_path in enumerate(rendered):
        image = cv2.imread(str(frame_path))
        image = cv2.resize(image, (panel_width, panel_height), interpolation=cv2.INTER_AREA)
        x = column * (panel_width + gutter)
        canvas[:panel_height, x : x + panel_width] = image
    figure = Image.fromarray(cv2.cvtColor(canvas, cv2.COLOR_BGR2RGB))
    draw = ImageDraw.Draw(figure)
    frame_font = ImageFont.truetype(FONT_PATH, 72)
    legend_font = ImageFont.truetype(FONT_PATH, 62)
    for column, frame_number in enumerate(FRAMES):
        draw.text((column * (panel_width + gutter) + 40, 32), f"#{frame_number:04d}", font=frame_font, fill="yellow")
    labels = [("Ground Truth", GT_COLOR), ("Visual Prompt", VP_COLOR), ("VPTracker", PREDICTION_COLOR)]
    for index, (label, color) in enumerate(labels):
        center_x = canvas_width * (index + 0.5) / len(labels)
        group_width = 180 + 34 + text_width(draw, label, legend_font)
        x = int(center_x - group_width / 2)
        draw.rectangle((x, panel_height + 45, x + 180, panel_height + 125), fill=tuple(reversed(color)))
        draw.text((x + 214, panel_height + 30), label, font=legend_font, fill="black")
    for path in FIGURE_PATHS:
        figure.save(path, compress_level=3)

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(4.45)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    panel_width_in, panel_height_in, gutter_in, y = 3.20, 2.13, 0.06, 0.10
    for column, frame_number in enumerate(FRAMES):
        index = frame_number - 1
        panel_x = 0.08 + column * (panel_width_in + gutter_in)
        frame_path = frames_root / f"frame_{frame_number:04d}{image_paths[index].suffix.lower()}"
        slide.shapes.add_picture(str(frame_path), Inches(panel_x), Inches(y), Inches(panel_width_in), Inches(panel_height_in))
        with Image.open(frame_path) as source:
            image_width, image_height = source.size
        prompt = vp_box(predictions[index - 1], image_width, image_height) if index else np.zeros(4)
        add_outline(slide, prompt, image_width, image_height, panel_x, y, panel_width_in, panel_height_in, VP_COLOR, 0.7)
        add_outline(slide, gt[index], image_width, image_height, panel_x, y, panel_width_in, panel_height_in, GT_COLOR, 1.5)
        add_outline(slide, predictions[index], image_width, image_height, panel_x, y, panel_width_in, panel_height_in, PREDICTION_COLOR, 1.5)
        add_text(slide, f"#{frame_number:04d}", panel_x + 0.05, y + 0.04, 0.65, 0.25, 17, RGBColor(255, 255, 0))
    labels = [("Ground Truth", GT_COLOR), ("Visual Prompt", VP_COLOR), ("VPTracker", PREDICTION_COLOR)]
    for index, (label, color) in enumerate(labels):
        center_x = 13.333 * (index + 0.5) / len(labels)
        swatch_x = center_x - 0.70
        swatch = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(swatch_x), Inches(2.70), Inches(0.48), Inches(0.22))
        swatch.fill.solid()
        swatch.fill.fore_color.rgb = rgb(color)
        swatch.line.fill.background()
        add_text(slide, label, swatch_x + 0.58, 2.64, 1.45, 0.30, 20)
    presentation.save(PPT_PATH)
    print(OUTPUT_ROOT)
    print(PPT_PATH)


if __name__ == "__main__":
    main()
